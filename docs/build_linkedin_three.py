"""
Build three LinkedIn-ready Word documents — distinct angles on the same
capstone project. Each is self-contained and copy-pastes cleanly into
LinkedIn's article editor.

Outputs:
    docs/linkedin_article_engineering.docx   (technical / architecture angle)
    docs/linkedin_article_product.docx       (product / UX angle)
    docs/linkedin_article_strategy.docx      (career / strategy angle)
"""
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT_DIR = Path(__file__).parent


# ── Helpers ────────────────────────────────────────────────────────────────────
def make_doc():
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    return doc


def h0(doc, text, size=24):
    p = doc.add_heading(text, level=0)
    for r in p.runs:
        r.font.size = Pt(size)
    return p


def h1(doc, text, size=16):
    p = doc.add_heading(text, level=1)
    for r in p.runs:
        r.font.size = Pt(size)
    return p


def body(doc, parts):
    """parts: list of (text, style_flags). flags in {'', 'b', 'i', 'bi'}."""
    p = doc.add_paragraph()
    for text, flags in parts:
        r = p.add_run(text)
        r.font.size = Pt(11)
        r.bold = "b" in flags
        r.italic = "i" in flags
    return p


def plain(doc, text):
    body(doc, [(text, "")])


def bullet(doc, parts, numbered=False):
    style = "List Number" if numbered else "List Bullet"
    p = doc.add_paragraph(style=style)
    for text, flags in parts:
        r = p.add_run(text)
        r.font.size = Pt(11)
        r.bold = "b" in flags
        r.italic = "i" in flags


def italic(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(10)
    r.italic = True


def divider(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("· · ·")
    r.font.size = Pt(11)


# ════════════════════════════════════════════════════════════════════════════
# ARTICLE 1 — Engineering / Architecture angle
# ════════════════════════════════════════════════════════════════════════════
def build_article_engineering():
    doc = make_doc()

    h0(doc, "Why I Built an AI Shopping Assistant That Never Lets the AI Rank a Product")
    italic(doc, "On the trust boundary between language work and decision logic")

    plain(doc,
        "I just shipped my MSBA capstone: a conversational shopping assistant that helps users "
        "move from a vague frustration ('my current bottle leaks in my gym bag') to a confident "
        "purchase decision. It uses Claude. But there is one place I deliberately kept the LLM "
        "out: ranking.")
    plain(doc, "That decision shaped everything else.")

    h1(doc, "The problem with letting an LLM rank")
    plain(doc,
        "A language model can absolutely rank products. Give it the catalogue, give it the "
        "user's preferences, and it will produce a top-N list. But three properties matter for "
        "decision support, and an LLM ranker fails all three:")
    bullet(doc, [("Reproducibility. ", "b"), ("Same input, different output across runs. Bad for audits, bad for debugging.", "")])
    bullet(doc, [("Stability across model upgrades. ", "b"), ("Today's GPT-X.Y top-1 is not next month's. Your 'best fit' shifts under your feet.", "")])
    bullet(doc, [("Explainability by construction. ", "b"), ("When the user asks 'why this and not that?', a post-hoc LLM rationale is just another generation. It can sound plausible without being grounded in the actual ranking signal.", "")])

    h1(doc, "The trust boundary")
    plain(doc, "I split the system into two halves with a hard line between them.")
    body(doc, [("Language side (LLM-eligible): ", "b"),
               ("parse free-text input ('phone screen too small for cooking' → category + use_case), "
                "map free-text chip answers, write personal per-product copy.", "")])
    body(doc, [("Decision side (deterministic only): ", "b"),
               ("filter by hard constraints, score with a two-tier weighted formula (general: rating · price · "
                "delivery; feature: category-specific rules like voice_control, insulated_gym, "
                "cabinet_stackable, leak_proof_match), produce three curated picks (Best Fit / Budget / "
                "Stretch), attach factual tradeoff labels (Cheapest, Most leakproof) — but only when "
                "factually true within the shown set.", "")])
    plain(doc,
        "The frontend's 'Why this fits' page maps each rule the ranker emitted to a human-readable "
        "sentence via a dictionary that lives on both sides. A reason can only appear on the page if "
        "its rule actually fired. That's grounded explanation by construction, not by hope.")

    h1(doc, "Every LLM call has a deterministic fallback")
    plain(doc, "The product runs end-to-end with ANTHROPIC_API_KEY=placeholder. If Claude is unreachable:")
    bullet(doc, [("Free-text input → a 30-pattern regex classifier picks the category ('water bottle', 'phone screen too small', 'pantry chaos').", "")])
    bullet(doc, [("Refine input ('under $80', 'larger screen', 'no camera') → a category-aware regex parser produces structured updates.", "")])
    bullet(doc, [("Per-product copy → a rule→text dictionary on the same code path the LLM uses.", "")])
    plain(doc,
        "This wasn't UX polish. It was an architectural test: ")
    body(doc, [("can the system survive the LLM layer being unreliable? ", "i"),
               ("If yes, the LLM is a contributor. If no, the LLM is a single point of failure dressed up as a feature.", "")])

    h1(doc, "The lesson")
    plain(doc,
        "The interesting question for AI-product builders is not 'where can we add an LLM?' — it is "
        "'where can we afford the answer to be probabilistic?' In a shopping decision, the language "
        "understanding can be probabilistic. The list of items I show you cannot.")
    plain(doc,
        "If you've shipped LLM-driven recommenders, I'd love to hear how you handle reproducibility "
        "and audits.")

    out = OUT_DIR / "linkedin_article_engineering.docx"
    doc.save(out)
    print(f"saved {out}")


# ════════════════════════════════════════════════════════════════════════════
# ARTICLE 2 — Product / UX angle
# ════════════════════════════════════════════════════════════════════════════
def build_article_product():
    doc = make_doc()

    h0(doc, "Most 'AI Shopping Assistants' Are Still Search Bars With a Chat Skin")
    italic(doc, "Five principles for building a needs-to-decision assistant instead")

    plain(doc,
        "I spent the last semester rebuilding what an AI shopping assistant should do. The thing I "
        "keep coming back to: shopping is a decision problem, not a search problem, and almost every "
        "tool on the market still treats it as the latter.")

    h1(doc, "The vague-need problem")
    plain(doc,
        "Real users don't start with 'Echo Show 8'. They start with 'I cook every day and my phone "
        "screen is too small to follow recipes while my hands are messy.' That sentence contains a "
        "category implication, three implicit preferences, and a use case. Today's assistants — "
        "Amazon Rufus, ChatGPT Shopping, Google Shopping, Perplexity — discard most of it. They take "
        "the keywords, run a search, and show you a list.")
    plain(doc, "The output looks like search results because it is search results.")

    h1(doc, "What a decision-oriented assistant does instead")
    plain(doc, "I built mine around five principles.")

    body(doc, [("1. Start from the frustration, not the category. ", "b"),
               ("Scene 1 asks 'what are you trying to solve?' — not 'what category?' If the user types "
                "'water bottle', the system politely says 'got it — let's narrow it down' and skips the "
                "empathy script. If they type the cooking complaint above, it infers phone screen + "
                "cooking and lands on smart displays.", "")])

    body(doc, [("2. Ask decision-relevant questions, not generic ones. ", "b"),
               ("For smart displays, that means voice ecosystem (Alexa/Google/Apple), placement "
                "(kitchen counter / wall / living room), screen size priority, privacy stance on "
                "cameras. Six chips, every one of them changes the ranking. Compare to 'what's your "
                "budget and how soon do you need it?' — which is exactly what every search filter "
                "already does.", "")])

    body(doc, [("3. Surface three picks, not twelve. ", "b"),
               ("A grid of similar-looking cards is not a decision aid; it is a comparison-shopping "
                "cop-out. I show Best Fit (highest match), Budget Pick (cheapest with ≥ 75% of best "
                "score), and a Stretch Pick (Largest Screen / Most Capacity / Best Visibility, "
                "depending on category). Each carries factual tradeoff labels — 'Cheapest', 'Most "
                "leakproof', 'Most portable' — only when actually true within the three shown.", "")])

    body(doc, [("4. Show what each pick honored AND what it missed. ", "b"),
               ("Every recommendation has a checklist:", "")])
    bullet(doc, [("✓ Material: stainless", "")])
    bullet(doc, [("✓ Insulated", "")])
    bullet(doc, [("✗ Large capacity (≥24oz) — ", ""), ("19oz", "i")])
    bullet(doc, [("? Drinking style: freesip — ", ""), ("unknown", "i")])
    plain(doc,
        "The trade-off becomes visible. The Budget Pick that's 19oz when the user wanted 24oz+ is "
        "not hidden — it is labeled. Trust is built by saying 'we couldn't satisfy this constraint' "
        "out loud.")

    body(doc, [("5. Compare 2–3 finalists side by side. ", "b"),
               ("A real comparison view, with category-aware rows (screen size for displays, "
                "capacity for bottles, visibility for organisers) and the row-winner highlighted. "
                "This was the single most-requested feature in walkthrough feedback.", "")])

    h1(doc, "The post-purchase part")
    plain(doc,
        "The lifecycle layer is the hardest one to get right. My first version implied the bottle "
        "could log refills. It can't — it's a piece of metal. The honest version says 'the bottle is "
        "dumb, the schedule is yours' and offers phone reminders, a gym-bag checklist, "
        "replacement-parts reorder. Less impressive. More useful.")

    h1(doc, "The thesis")
    plain(doc,
        "Existing tools assume the user already knows what they want. The strongest version of an AI "
        "assistant assumes they don't, and helps them figure it out without pretending to be magic.")
    plain(doc,
        "If you're building anything in commerce + AI, I'd be curious whether your product currently "
        "behaves like a search bar or like a decision partner.")

    out = OUT_DIR / "linkedin_article_product.docx"
    doc.save(out)
    print(f"saved {out}")


# ════════════════════════════════════════════════════════════════════════════
# ARTICLE 3 — Career / strategy angle
# ════════════════════════════════════════════════════════════════════════════
def build_article_strategy():
    doc = make_doc()

    h0(doc, "What I Learned by Building Things My AI Was NOT Allowed to Do")
    italic(doc, "Disciplined AI use, told through a 12-week capstone")

    plain(doc,
        "Halfway through my MSBA capstone, I realised I was learning more about AI by carefully "
        "limiting where it could touch the product than by adding more of it.")

    h1(doc, "The default failure mode")
    plain(doc,
        "The temptation in any 12-week applied AI project is to wrap an LLM around every step. "
        "Vague input? Throw it at GPT. Need a recommendation? Throw it at GPT. Trade-off "
        "explanation? GPT again. The result is a product that looks impressive in a demo and falls "
        "apart on inspection: same input gives different output, 'best' picks shift between runs, "
        "explanations are confidently wrong, and the whole thing collapses if the API has an outage.")
    plain(doc, "I went the other way.")

    h1(doc, "The discipline I imposed")
    plain(doc,
        "For every spot where I could have called an LLM, I asked: ")
    body(doc, [("would I trust this answer in production? ", "i"),
               ("If the answer is 'yes, even if it's slightly different each time' — Claude was "
                "welcome. If the answer is 'no, I need this to be reproducible and auditable' — "
                "Claude was banned.", "")])
    plain(doc, "The split came out roughly:")
    bullet(doc, [("Use Claude for: ", "b"),
                 ("parsing free-text intent, mapping free-text chip answers to structured values, "
                  "writing personal one-line product copy.", "")])
    bullet(doc, [("Don't use Claude for: ", "b"),
                 ("filtering by hard constraints, scoring/ranking products, deciding which trade-offs "
                  "to surface, generating explanations not grounded in the ranker's actual signals.", "")])
    plain(doc,
        "Then I added a deterministic fallback for every LLM call, so the entire product runs "
        "without an API key. That fallback path isn't a 'graceful degradation' hack — it's the "
        "architectural commitment that the system's correctness doesn't depend on the AI being "
        "available.")

    h1(doc, "Three things this taught me")

    body(doc, [("1. Grounded explanations are dramatically more trustworthy than post-hoc rationales. ", "b"),
               ("I built a small dictionary that maps each scoring rule (voice_control, insulated_gym, "
                "cabinet_stackable) to a one-line user-facing explanation. The rule fires in the "
                "ranker, the explanation appears on the page. There is no LLM step in between to "
                "invent reasons that sound plausible but aren't actually why the product was picked. "
                "A 26-row Python dict beat a 200B-parameter LLM at the explainability job.", "")])

    body(doc, [("2. Honesty about what's automated builds trust faster than feature claims. ", "b"),
               ("My first lifecycle screen for water bottles said 'track every refill, get nudged "
                "when you fall behind'. Sounded great. Was misleading — a regular bottle doesn't "
                "sense anything. The honest version says 'the bottle is dumb; set 2-hour phone "
                "reminders'. It tested better. Users gave more credit to a system that admitted its "
                "limits than one that pretended to be smarter than it was.", "")])

    body(doc, [("3. The most defensible AI products treat AI as a bounded contributor, not the core. ", "b"),
               ("When I ask a hiring manager what they want from an analytics hire, 'knows when to "
                "use AI' is more interesting than 'knows how to call OpenAI'. Anyone can call "
                "OpenAI. The skill is figuring out where the LLM helps and where it actively makes "
                "the product worse.", "")])

    h1(doc, "Where I'd apply this beyond shopping")
    bullet(doc, [("Hiring tools: ", "b"),
                 ("LLM for parsing résumé free text, deterministic ranker for the actual scoring.", "")])
    bullet(doc, [("Healthcare triage: ", "b"),
                 ("LLM for capturing patient narrative, rule-based decision support for triage tier.", "")])
    bullet(doc, [("Financial planning: ", "b"),
                 ("LLM for goal elicitation, deterministic optimisation for portfolio construction.", "")])
    plain(doc, "In all three, the language model is the interview, not the doctor.")

    h1(doc, "The takeaway")
    plain(doc,
        "The capstone shipped, but the takeaway is portable: the most disciplined AI use case is the "
        "one where you can confidently say 'the LLM is not allowed here, and here is why.'")
    plain(doc,
        "If you're hiring or building in applied AI, I'd be glad to talk about where the boundaries "
        "belong.")

    out = OUT_DIR / "linkedin_article_strategy.docx"
    doc.save(out)
    print(f"saved {out}")


if __name__ == "__main__":
    build_article_engineering()
    build_article_product()
    build_article_strategy()
